# -*- coding: utf-8 -*-
# =============================================================
#  TEMPLATE LAPORAN (reusable) - gaya sesuai preferensi Irfan
#  Font Times New Roman, heading hitam, cover: Nama/NIM/Kelas.
#
#  Cara pakai:
#     from laporan_template import build_laporan
#     build_laporan(
#         pdf_path="...pdf", docx_path="...docx" or None,
#         title="TRANSFORMASI DATA (NORMALISASI)",
#         subtitle="Mata Kuliah Machine Learning – Pertemuan 3",
#         footer_slug="434241079_Normalisasi",
#         blocks=[("h1","1. Pendahuluan"), ("p","..."),
#                 ("table",(hdr, rows)), ("img",(path,16.0)), ("cap","..."),
#                 ("code","..."), ("term","..."), ("pb",None)])
#
#  Identitas default diambil dari IDENTITAS di bawah (boleh dioverride).
# =============================================================
import os, html, re

IDENTITAS = {
    "nama":  "Muhammad Irfan Nuha",
    "nim":   "434241079",
    "kelas": "TI-C5",
    "status":"Tugas Individu",
}

# ---- ukuran (pt) ----
SZ_PRE, SZ_TITLE, SZ_SUB = 18, 24, 13
SZ_ID, SZ_H1, SZ_H2, SZ_BODY, SZ_CAP = 13, 16, 13, 12, 9
INK = "#000000"          # heading & teks hitam (Times New Roman, akademik)


# =================================================================
#  PDF (reportlab) - Times New Roman = Times-Roman bawaan
# =================================================================
def _build_pdf(path, title, subtitle, blocks, ident, footer_slug):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Image,
                                     Table, TableStyle, PageBreak, HRFlowable)
    BLACK = colors.HexColor(INK)
    H1 = ParagraphStyle("H1", fontName="Times-Bold", fontSize=SZ_H1, textColor=BLACK, spaceBefore=8, spaceAfter=6, leading=SZ_H1*1.2)
    H2 = ParagraphStyle("H2", fontName="Times-Bold", fontSize=SZ_H2, textColor=BLACK, spaceBefore=6, spaceAfter=4, leading=SZ_H2*1.2)
    BODY = ParagraphStyle("BODY", fontName="Times-Roman", fontSize=SZ_BODY, leading=SZ_BODY*1.35, alignment=TA_JUSTIFY)
    CAP = ParagraphStyle("CAP", fontName="Times-Italic", fontSize=SZ_CAP, textColor=colors.grey, alignment=TA_CENTER, spaceBefore=2, spaceAfter=10)
    CELL = ParagraphStyle("CELL", fontName="Times-Roman", fontSize=8.2, leading=10)
    HEADC = ParagraphStyle("HEADC", fontName="Times-Bold", fontSize=8.2, leading=10, textColor=colors.white)
    CODE = ParagraphStyle("CODE", fontName="Courier", fontSize=8.5, leading=11, textColor=colors.HexColor("#0b3d0b"))
    TERM = ParagraphStyle("TERM", fontName="Courier", fontSize=8, leading=10.5, textColor=colors.whitesmoke)

    def md(s):
        s = html.escape(s); s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
        return s.replace("\n", "<br/>")

    def table(hdr, rows):
        data = [[Paragraph(md(str(h)), HEADC) for h in hdr]] + \
               [[Paragraph(md(str(c)), CELL) for c in r] for r in rows]
        ncol = len(hdr); w0 = 4.5*cm if ncol == 5 else (16.6*cm/ncol)
        rest = (16.6*cm - w0)/(ncol-1) if ncol > 1 else w0
        t = Table(data, colWidths=[w0] + [rest]*(ncol-1), repeatRows=1)
        st = [("BACKGROUND",(0,0),(-1,0),colors.HexColor("#404040")),
              ("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#b8b8b8")),
              ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
              ("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3),
              ("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4)]
        for i in range(1, len(data)):
            if i % 2 == 0: st.append(("BACKGROUND",(0,i),(-1,i),colors.HexColor("#f0f0f0")))
        t.setStyle(TableStyle(st)); return t

    def boxpara(text, style, bg, border):
        p = Paragraph(html.escape(text).replace(" ","&nbsp;").replace("\n","<br/>"), style)
        tt = Table([[p]], colWidths=[16.6*cm])
        tt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),bg),("BOX",(0,0),(-1,-1),0.5,border),
            ("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6),
            ("LEFTPADDING",(0,0),(-1,-1),8),("RIGHTPADDING",(0,0),(-1,-1),8)])); return tt

    def image(pw):
        p, w = pw; iw, ih = ImageReader(p).getSize()
        natural_cm = iw/130*2.54
        W = min(w, natural_cm)*cm
        return Image(p, width=W, height=W*ih/iw)

    story = []
    # ---- COVER ----
    C = lambda sz, b=False: ParagraphStyle("c", fontName="Times-Bold" if b else "Times-Roman",
                                           fontSize=sz, alignment=TA_CENTER, leading=sz*1.3)
    story += [Spacer(1, 3.2*cm),
        Paragraph("LAPORAN TUGAS", C(SZ_PRE, True)),
        Paragraph(title, C(SZ_TITLE, True)),
        Spacer(1, 0.3*cm),
        Paragraph(subtitle, C(SZ_SUB)),
        Spacer(1, 2.2*cm),
        HRFlowable(width="55%", thickness=1, color=colors.HexColor("#555555"), hAlign="CENTER"),
        Spacer(1, 0.5*cm),
        Paragraph(f"NAMA : {ident['nama']}", C(SZ_ID, True)),
        Paragraph(f"NIM : {ident['nim']}", C(SZ_ID, True)),
        Paragraph(f"Kelas : {ident['kelas']}", C(SZ_ID, True)),
        Spacer(1, 0.2*cm),
        Paragraph(ident["status"], C(SZ_SUB)),
        Spacer(1, 0.5*cm),
        HRFlowable(width="55%", thickness=1, color=colors.HexColor("#555555"), hAlign="CENTER"),
        PageBreak()]

    for kind, val in blocks:
        if kind == "h1": story.append(Paragraph(md(val), H1))
        elif kind == "h2": story.append(Paragraph(md(val), H2))
        elif kind == "p": story += [Paragraph(md(val), BODY), Spacer(1, 0.12*cm)]
        elif kind == "cap": story.append(Paragraph(md(val), CAP))
        elif kind == "table": story += [table(*val), Spacer(1, 0.05*cm)]
        elif kind == "code": story += [boxpara(val, CODE, colors.HexColor("#f2f6ec"), colors.HexColor("#b8c9a8")), Spacer(1,0.15*cm)]
        elif kind == "term": story += [boxpara(val, TERM, colors.HexColor("#1e1e1e"), colors.HexColor("#444")), Spacer(1,0.1*cm)]
        elif kind == "img": story.append(image(val))
        elif kind == "pb": story.append(PageBreak())

    def foot(c, doc):
        c.saveState(); c.setFont("Times-Roman", 8); c.setFillColor(colors.grey)
        c.drawString(2*cm, 1.1*cm, footer_slug)
        c.drawRightString(19*cm, 1.1*cm, f"Halaman {doc.page}"); c.restoreState()

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=1.8*cm, bottomMargin=1.6*cm, title=footer_slug, author=ident["nim"])
    doc.build(story, onFirstPage=foot, onLaterPages=foot)
    print("PDF :", path)


# =================================================================
#  DOCX (python-docx) - Times New Roman
# =================================================================
def _build_docx(path, title, subtitle, blocks, ident, footer_slug):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from PIL import Image as PILImage
    BLACK = RGBColor(0, 0, 0)

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.9); s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.9); s.right_margin = Inches(0.9)
    nf = doc.styles["Normal"].font; nf.name = "Times New Roman"; nf.size = Pt(SZ_BODY)

    def shade(cell, hexc):
        tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexc); tcPr.append(sh)

    def tnr(run, size=None, bold=None, color=None):
        run.font.name = "Times New Roman"
        if size: run.font.size = Pt(size)
        if bold is not None: run.bold = bold
        if color is not None: run.font.color.rgb = color

    def add_runs(p, text):
        for part in re.split(r"(\*\*.+?\*\*)", text):
            if not part: continue
            if part.startswith("**") and part.endswith("**"):
                tnr(p.add_run(part[2:-2]), bold=True)
            else:
                tnr(p.add_run(part))

    def center(txt, size, bold=False):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tnr(p.add_run(txt), size=size, bold=bold); return p

    def heading(txt, size):
        h = doc.add_heading(level=1 if size >= SZ_H1 else 2)
        h.paragraph_format.space_before = Pt(6)
        for r in h.runs: r.text = ""
        tnr(h.add_run(txt), size=size, bold=True, color=BLACK)

    def add_table(hdr, rows):
        t = doc.add_table(rows=1, cols=len(hdr)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
        for j, h in enumerate(hdr):
            c = t.rows[0].cells[j]; shade(c, "404040")
            tnr(c.paragraphs[0].add_run(str(h)), size=8, bold=True, color=RGBColor(255,255,255))
        for i, row in enumerate(rows):
            cells = t.add_row().cells
            for j, v in enumerate(row):
                if i % 2 == 1: shade(cells[j], "F0F0F0")
                tnr(cells[j].paragraphs[0].add_run(str(v)), size=8)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def mono(text, dark):
        tbl = doc.add_table(rows=1, cols=1); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0); shade(cell, "1E1E1E" if dark else "F2F6EC")
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        for i, line in enumerate(text.split("\n")):
            if i: p.add_run().add_break()
            r = p.add_run(line if line else " ")
            r.font.name = "Consolas"; r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0xF2,0xF2,0xF2) if dark else RGBColor(0x0b,0x3d,0x0b)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---- COVER ----
    for _ in range(6): doc.add_paragraph()
    center("LAPORAN TUGAS", SZ_PRE, True)
    center(title, SZ_TITLE, True)
    center(subtitle, SZ_SUB)
    for _ in range(3): doc.add_paragraph()
    center(f"NAMA : {ident['nama']}", SZ_ID, True)
    center(f"NIM : {ident['nim']}", SZ_ID, True)
    center(f"Kelas : {ident['kelas']}", SZ_ID, True)
    center(ident["status"], SZ_SUB)
    doc.add_page_break()

    for kind, val in blocks:
        if kind == "h1": heading(val, SZ_H1)
        elif kind == "h2": heading(val, SZ_H2)
        elif kind == "p":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(8); add_runs(p, val)
        elif kind == "cap":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tnr(p.add_run(val), size=SZ_CAP, color=RGBColor(0x80,0x80,0x80)); p.runs[0].italic = True
        elif kind == "table": add_table(*val)
        elif kind == "code": mono(val, dark=False)
        elif kind == "term": mono(val, dark=True)
        elif kind == "img":
            p, w = val; iw, _ = PILImage.open(p).size; natural_in = iw/130
            par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(p, width=Inches(min(w/2.54, natural_in)))
        elif kind == "pb": doc.add_page_break()

    ft = doc.sections[0].footer.paragraphs[0]; ft.text = ""
    tnr(ft.add_run(footer_slug), size=8, color=RGBColor(0x88,0x88,0x88))
    doc.save(path); print("DOCX:", path)


def build_laporan(pdf_path, title, subtitle, blocks, footer_slug,
                  docx_path=None, ident=None):
    ident = {**IDENTITAS, **(ident or {})}
    if pdf_path:  _build_pdf(pdf_path, title, subtitle, blocks, ident, footer_slug)
    if docx_path: _build_docx(docx_path, title, subtitle, blocks, ident, footer_slug)
